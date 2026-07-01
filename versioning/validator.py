import os
import docx

class DocumentValidator:
    @staticmethod
    def validate_docx(file_path: str) -> bool:
        """
        Runs sanity checks on the docx file.
        Returns True if valid, raises ValueError if issues are found.
        """
        if not os.path.exists(file_path):
            raise ValueError(f"File not found: {file_path}")

        try:
            doc = docx.Document(file_path)
        except Exception as e:
            raise ValueError(f"Failed to parse Word document: {str(e)}")

        # 1. Check if document is empty
        if len(doc.paragraphs) == 0 and len(doc.tables) == 0:
            raise ValueError("Document is empty (no paragraphs or tables).")

        # 2. Check if heading structure is present (must have at least one Heading 1)
        has_headings = False
        for p in doc.paragraphs:
            if p.style and p.style.name and p.style.name.startswith("Heading 1"):
                has_headings = True
                break
                
        if not has_headings:
            raise ValueError("Document violates hierarchical schema: No 'Heading 1' styles detected.")

        return True

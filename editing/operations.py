import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from typing import Dict, Any, List
from editing.locator import DocumentLocator

def set_cell_borders(cell, color_hex: str = "000000", sz: str = "12"):
    """Applies borders to a single cell via XML manipulation."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    
    for border_name in ['top', 'left', 'bottom', 'right']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), sz)  # 12 is 1.5 pt
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), color_hex)
        tcBorders.append(border)
        
    tcPr.append(tcBorders)

def safe_replace_text(paragraph, new_text: str):
    """Replaces text in a paragraph while preserving paragraph and run formatting."""
    if paragraph.runs:
        # Update first run text
        paragraph.runs[0].text = new_text
        # Clear text in subsequent runs so we don't duplicate
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(new_text)

class ReplaceText:
    def __init__(self, paragraph_id: int, new_text: str, section: str, description: str):
        self.paragraph_id = paragraph_id
        self.new_text = new_text
        self.section = section
        self.description = description

    def apply(self, doc: docx.document.Document):
        p = DocumentLocator.find_paragraph_by_id(doc, self.paragraph_id)
        if not p:
            raise ValueError(f"Paragraph ID {self.paragraph_id} not found.")
        safe_replace_text(p, self.new_text)

class InsertWarning:
    def __init__(self, after_paragraph_id: int, warning_text: str, section: str, description: str):
        self.after_paragraph_id = after_paragraph_id
        self.warning_text = warning_text
        self.section = section
        self.description = description

    def apply(self, doc: docx.document.Document):
        p = DocumentLocator.find_paragraph_by_id(doc, self.after_paragraph_id)
        if not p:
            raise ValueError(f"Paragraph ID {self.after_paragraph_id} not found.")
        
        # Create a single-cell table
        table = doc.add_table(rows=1, cols=1)
        # Move it in the XML tree directly after paragraph
        p._p.addnext(table._tbl)
        
        # Add warning text
        cell = table.rows[0].cells[0]
        set_cell_borders(cell, color_hex="FF0000", sz="12") # Red borders for WARNING
        
        cell_p = cell.paragraphs[0]
        if "Normal" in doc.styles:
            cell_p.style = doc.styles["Normal"]
        # Format text prefix as WARNING
        run_bold = cell_p.add_run("WARNING: ")
        run_bold.bold = True
        cell_p.add_run(self.warning_text)

class InsertCaution:
    def __init__(self, after_paragraph_id: int, caution_text: str, section: str, description: str):
        self.after_paragraph_id = after_paragraph_id
        self.caution_text = caution_text
        self.section = section
        self.description = description

    def apply(self, doc: docx.document.Document):
        p = DocumentLocator.find_paragraph_by_id(doc, self.after_paragraph_id)
        if not p:
            raise ValueError(f"Paragraph ID {self.after_paragraph_id} not found.")
        
        # Create single-cell table
        table = doc.add_table(rows=1, cols=1)
        p._p.addnext(table._tbl)
        
        # Add caution text
        cell = table.rows[0].cells[0]
        set_cell_borders(cell, color_hex="FF8C00", sz="12") # Amber borders for CAUTION
        
        cell_p = cell.paragraphs[0]
        if "Normal" in doc.styles:
            cell_p.style = doc.styles["Normal"]
        run_bold = cell_p.add_run("CAUTION: ")
        run_bold.bold = True
        cell_p.add_run(self.caution_text)

class ReplaceStep:
    def __init__(self, section: str, step_number: str, new_text: str, description: str):
        self.section = section
        self.step_number = step_number
        self.new_text = new_text
        self.description = description

    def apply(self, doc: docx.document.Document):
        p = DocumentLocator.find_step_paragraph(doc, self.section, self.step_number)
        if not p:
            raise ValueError(f"Step {self.step_number} in section '{self.section}' not found.")
        
        # Re-write step prefix + text
        full_text = f"{self.step_number}. {self.new_text}"
        safe_replace_text(p, full_text)

class AppendSubsection:
    def __init__(self, parent_section: str, heading: str, content: str, description: str):
        self.parent_section = parent_section
        self.heading = heading
        self.content = content
        self.description = description

    def apply(self, doc: docx.document.Document):
        last_p, last_idx = DocumentLocator.find_section_end_paragraph(doc, self.parent_section)
        if not last_p:
            raise ValueError(f"Parent section '{self.parent_section}' not found.")
        
        # We need to insert a Heading 3 paragraph and body content paragraph after last_p
        # XML-wise insertion helper
        def insert_paragraph_after(target_p, text: str, style_name: str) -> docx.text.paragraph.Paragraph:
            new_p_element = OxmlElement('w:p')
            target_p._p.addnext(new_p_element)
            new_p = docx.text.paragraph.Paragraph(new_p_element, doc)
            new_p.text = text
            if style_name in doc.styles:
                new_p.style = doc.styles[style_name]
            return new_p

        # Insert body paragraph first, then heading, so that heading comes before body in the XML sequence
        # (Since we are doing addnext, calling it on last_p inserts directly after last_p, pushing previous addnext down!)
        # So:
        # last_p -> addnext(body) results in: last_p -> body
        # last_p -> addnext(heading) results in: last_p -> heading -> body
        # This is correct!
        body_p = insert_paragraph_after(last_p, self.content, "Normal")
        heading_p = insert_paragraph_after(last_p, self.heading, "Heading 3")

class UpdateTableCell:
    def __init__(self, table_id: int, row: int, col: int, new_value: str, section: str, description: str):
        self.table_id = table_id
        self.row = row
        self.col = col
        self.new_value = new_value
        self.section = section
        self.description = description

    def apply(self, doc: docx.document.Document):
        t = DocumentLocator.find_table_by_id(doc, self.table_id)
        if not t:
            raise ValueError(f"Table ID {self.table_id} not found.")
        
        if self.row >= len(t.rows) or self.col >= len(t.columns):
            raise ValueError(f"Cell coordinates ({self.row}, {self.col}) out of bounds for table {self.table_id}.")
            
        cell = t.rows[self.row].cells[self.col]
        # Preserving cell paragraph/runs styling if possible
        if cell.paragraphs:
            safe_replace_text(cell.paragraphs[0], self.new_value)
        else:
            cell.add_paragraph(self.new_value)

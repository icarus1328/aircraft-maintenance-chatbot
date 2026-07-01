import docx
from editing.locator import DocumentLocator

class EditValidator:
    @staticmethod
    def validate_replace_text(docx_path: str, paragraph_id: int) -> bool:
        doc = docx.Document(docx_path)
        p = DocumentLocator.find_paragraph_by_id(doc, paragraph_id)
        if not p:
            raise ValueError(f"Target paragraph with ID {paragraph_id} not found in document.")
        return True

    @staticmethod
    def validate_insert_warning(docx_path: str, after_paragraph_id: int) -> bool:
        doc = docx.Document(docx_path)
        p = DocumentLocator.find_paragraph_by_id(doc, after_paragraph_id)
        if not p:
            raise ValueError(f"Target paragraph with ID {after_paragraph_id} not found in document.")
        return True

    @staticmethod
    def validate_replace_step(docx_path: str, section: str, step_number: str) -> bool:
        doc = docx.Document(docx_path)
        p = DocumentLocator.find_step_paragraph(doc, section, step_number)
        if not p:
            raise ValueError(f"Step {step_number} in section '{section}' not found in document.")
        return True

    @staticmethod
    def validate_append_subsection(docx_path: str, parent_section: str) -> bool:
        doc = docx.Document(docx_path)
        last_p, _ = DocumentLocator.find_section_end_paragraph(doc, parent_section)
        if not last_p:
            raise ValueError(f"Parent section '{parent_section}' not found in document.")
        return True

    @staticmethod
    def validate_update_table_cell(docx_path: str, table_id: int, row: int, col: int) -> bool:
        doc = docx.Document(docx_path)
        t = DocumentLocator.find_table_by_id(doc, table_id)
        if not t:
            raise ValueError(f"Table with ID {table_id} not found in document.")
        if row >= len(t.rows) or col >= len(t.columns):
            raise ValueError(f"Cell coordinates ({row}, {col}) out of bounds for table {table_id} (dimensions: {len(t.rows)}x{len(t.columns)}).")
        return True
